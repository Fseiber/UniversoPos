"""Lee el Excel sin modificarlo y prepara registros con trazabilidad por fila."""
import argparse
import hashlib
import json
import math
from pathlib import Path
from decimal import Decimal
import openpyxl

parser = argparse.ArgumentParser()
parser.add_argument("source")
parser.add_argument("output")
args = parser.parse_args()
source = Path(args.source)
workbook = openpyxl.load_workbook(source, read_only=True, data_only=False)
try:
    if workbook.sheetnames != ["Catálogo"]:
        raise ValueError("La estructura del archivo cambió: revisar las hojas antes de importar")
    sheet = workbook["Catálogo"]
    rows = sheet.iter_rows(values_only=True)
    if tuple(next(rows)) != ("Nombre del Producto", "Precio (ARS)", "Unidad de Medida"):
        raise ValueError("Las columnas no coinciden con el catálogo revisado")
    records = []
    seen = set()
    for row_number, row in enumerate(rows, 2):
        if all(value is None for value in row):
            continue
        name, price, unit = row
        if not isinstance(name, str) or not name.strip() or len(name.strip()) > 120:
            raise ValueError(f"Nombre inválido en fila {row_number}")
        if isinstance(price, bool) or not isinstance(price, (int, float)) or not math.isfinite(price) or price < 0:
            raise ValueError(f"Precio de referencia inválido en fila {row_number}")
        amount = Decimal(str(price))
        if amount != amount.quantize(Decimal("0.01")) or amount > 100_000_000:
            raise ValueError(f"Precisión de precio inválida en fila {row_number}")
        units = {"UNIDAD": "UNIDAD", "KILO": "PESO"}
        if unit not in units:
            raise ValueError(f"Unidad desconocida en fila {row_number}: {unit}")
        key = (" ".join(name.strip().casefold().split()), units[unit])
        if key in seen:
            raise ValueError(f"Nombre y unidad duplicados en fila {row_number}")
        seen.add(key)
        records.append({"sourceRow": row_number, "name": name.strip(), "unit": units[unit], "referencePrice": str(amount)})
    payload = {"sourceName": source.name, "sourceHash": hashlib.sha256(source.read_bytes()).hexdigest(), "sheetName": sheet.title, "records": records}
    target = Path(args.output)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"records": len(records), "byUnit": {unit: sum(r["unit"] == unit for r in records) for unit in units.values()}, "sourceHash": payload["sourceHash"], "output": str(target)}, ensure_ascii=True))
finally:
    workbook.close()
